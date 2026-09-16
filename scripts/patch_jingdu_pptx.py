#!/usr/bin/env python3
"""Patch 组会分享_AlphaCast_精读.pptx: paper fact-fixes + bounded repro slides."""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "docs" / "组会分享_AlphaCast_精读.pptx"


def _set_shape_text(shape, new_text: str) -> None:
    tf = shape.text_frame
    paras = list(tf.paragraphs)
    lines = new_text.split("\n")
    while len(paras) < len(lines):
        tf.add_paragraph()
        paras = list(tf.paragraphs)
    for i, para in enumerate(paras):
        line = lines[i] if i < len(lines) else ""
        if para.runs:
            para.runs[0].text = line
            for run in para.runs[1:]:
                run.text = ""
        else:
            run = para.add_run()
            run.text = line
    for i in range(len(lines), len(paras)):
        if paras[i].runs:
            paras[i].runs[0].text = ""
            for run in paras[i].runs[1:]:
                run.text = ""


def _replace_in_slide(slide, mapping: dict[str, str]) -> int:
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        old = sh.text_frame.text
        new = old
        for a, b in mapping.items():
            if a in new:
                new = new.replace(a, b)
        if new != old:
            _set_shape_text(sh, new)
            n += 1
    try:
        notes = slide.notes_slide.notes_text_frame
        old = notes.text
        new = old
        for a, b in mapping.items():
            if a in new:
                new = new.replace(a, b)
        if new != old:
            notes.text = new
            n += 1
    except Exception:
        pass
    return n


def _duplicate_slide(prs: Presentation, index: int) -> None:
    source = list(prs.slides)[index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shape in list(dest.shapes):
        shape.element.getparent().remove(shape.element)

    rid_map: dict[str, str] = {}
    for rel in source.part.rels.values():
        if "image" in rel.reltype:
            new_rid = dest.part.relate_to(rel.target_part, rel.reltype)
            rid_map[rel.rId] = new_rid

    for shape in source.shapes:
        new_el = copy.deepcopy(shape.element)
        for blip in new_el.xpath(".//*[local-name()='blip']"):
            embed = blip.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
            )
            if embed in rid_map:
                blip.set(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
                    rid_map[embed],
                )
        dest.shapes._spTree.insert_element_before(new_el, "p:extLst")


def _reorder_slides(prs: Presentation, order: list[int]) -> None:
    sld_id_lst = prs.slides._sldIdLst
    entries = list(sld_id_lst)
    if len(entries) != len(order):
        raise RuntimeError(f"reorder size mismatch: {len(entries)} vs {len(order)}")
    new_entries = [entries[i] for i in order]
    for el in list(sld_id_lst):
        sld_id_lst.remove(el)
    for el in new_entries:
        sld_id_lst.append(el)


def _clear_slide_text_shapes(slide) -> None:
    """Remove non-background text/picture clutter; keep full-bleed bg if any."""
    for sh in list(slide.shapes):
        # keep very large background autoshape
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            if sh.width and sh.height and sh.width >= 11_000_000 and sh.height >= 6_000_000:
                continue
        sh.element.getparent().remove(sh.element)


def _add_textbox(slide, left, top, width, height, text: str, *, size=18, bold=False, color=(26, 46, 40)):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)
        try:
            run.font.name = "Microsoft YaHei"
        except Exception:
            pass
    return box


def _fill_repro_progress(slide) -> None:
    _clear_slide_text_shapes(slide)
    _add_textbox(slide, 762000, 400000, 10668000, 600000, "本机复现 · 进度与对齐边界", size=28, bold=True)
    _add_textbox(
        slide,
        762000,
        1050000,
        10668000,
        500000,
        "CastMind：在能对齐的维度对齐；不能对齐的写清残差。本页不是论文 Table 1。",
        size=16,
        color=(90, 107, 100),
    )
    body = (
        "对齐：协议（ETTh1 · L/H=96 · MSE/MAE）· 方法集命名映射 · 统计栈方法\n"
        "量级可对：Chronos / Sundial 公开权重（设备差为次要残差）\n"
        "近似：深度学习自训（无作者主表 .pth）——不可文件对齐\n"
        "不可数值对齐：系统行 DeepSeek ≠ 论文 GPT-5；不宣称复现 Table 1 系统格\n"
        "对照表：可做 Table 1 版式汇编；本机格 ≠ 论文格"
    )
    _add_textbox(slide, 762000, 1650000, 10668000, 3800000, body, size=18)
    _add_textbox(
        slide,
        762000,
        5600000,
        10668000,
        700000,
        "页脚：本页描述本机 CastMind 工作边界，数字若出现须标「本机、非论文」。",
        size=14,
        color=(122, 90, 32),
    )
    _add_textbox(slide, 762000, 6400000, 8000000, 300000, "AlphaCast · 论文精读与组会分享", size=12, color=(120, 120, 120))
    slide.notes_slide.notes_text_frame.text = (
        "本页讲复现边界，不报本机 MSE 冒充 Table 1。"
        "协议与方法集可对齐；DL 无官方权重；系统骨干 DeepSeek 不可数值对齐论文系统行。"
        "素材见 PAPER_ALIGNMENT.md。"
    )


def _fill_repro_issues(slide) -> None:
    _clear_slide_text_shapes(slide)
    _add_textbox(slide, 762000, 400000, 10668000, 600000, "本机复现 · 遇到的问题", size=28, bold=True)
    _add_textbox(
        slide,
        762000,
        1050000,
        10668000,
        500000,
        "开放项与卡点（事实陈述；不把本机结果写成论文结论）",
        size=16,
        color=(90, 107, 100),
    )
    body = (
        "1. 作者主表 DL 权重未公开 → 只能自训近似，不可宣称格级一致\n"
        "2. 系统 LLM：论文 GPT-5，本机 DeepSeek → 流程可跑通，数值不可对齐系统行\n"
        "3. 知识库 / 案例库构建与维护细节论文未写全 → 工程与复现边界\n"
        "4. 多轮智能体成本 / 延迟：论文未充分量化；本机可观察耗时但不作论文对照\n"
        "5. 部分数据域（如 Windy / Sunny / MOPEX）存在获取缺口时，不编造结果"
    )
    _add_textbox(slide, 762000, 1650000, 10668000, 3800000, body, size=18)
    _add_textbox(
        slide,
        762000,
        5600000,
        10668000,
        700000,
        "口径：能对齐的说对齐；不能的说不能。不问「是否完全复现成功」的二元结论。",
        size=14,
        color=(122, 90, 32),
    )
    _add_textbox(slide, 762000, 6400000, 8000000, 300000, "AlphaCast · 论文精读与组会分享", size=12, color=(120, 120, 120))
    slide.notes_slide.notes_text_frame.text = (
        "卡点五条：无官方 DL 权重、骨干不一致、知识/案例细节不全、成本未量化、数据缺口。"
        "被问「复现了吗」：回答对齐等级，不回答假的 Table 1 数字。"
    )


def _update_page_numbers(prs: Presentation) -> None:
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides):
        want = f"{i + 1:02d} / {total:02d}"
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            # match patterns like 14 / 20 or 02 / 20
            if " / " in t and len(t) <= 10 and t.replace(" ", "").replace("/", "").isdigit():
                _set_shape_text(sh, want)
            elif t.endswith(f" / {total}") or (len(t) <= 10 and " / 20" in t) or (len(t) <= 10 and " / 22" in t):
                _set_shape_text(sh, want)


def patch() -> None:
    if not PPTX.is_file():
        raise SystemExit(f"missing {PPTX}")
    prs = Presentation(str(PPTX))

    # --- paper fact fixes ---
    _replace_in_slide(
        prs.slides[0],
        {
            "所有实验数字以论文 PDF 为准": (
                "论文区数字均摘自论文表/图；复现页另计、不冒充 Table 1"
            ),
            "预计 25 分钟": "预计 28–30 分钟",
        },
    )
    _replace_in_slide(
        prs.slides[1],
        {
            "汇报路线 · 本次汇报的五个问题": "汇报路线 · 六个问题（含本机复现边界）",
            "价值判断 · 局限 · 开放问题": "价值判断 · 局限 · 本机复现边界",
        },
    )
    # Expand slide 2 item 05 area is cramped; patch notes instead of forcing a 06 card
    try:
        n = prs.slides[1].notes_slide.notes_text_frame
        if "复现" not in n.text:
            n.text = (
                n.text
                + "\n【补充】论文结果讲完后，用 1–2 页说明本机 CastMind 对齐等级与卡点；"
                "本机数字不与论文 Table 1 混谈。"
            )
    except Exception:
        pass

    _replace_in_slide(
        prs.slides[3],
        {"传统流程完全没有利用这些信息": "多数仍一次性映射，缺少显式的多源取证与反思"},
    )
    # soften notes if any
    try:
        nt = prs.slides[3].notes_slide.notes_text_frame
        nt.text = nt.text.replace("完全没有", "多数缺少显式利用与反思机制")
    except Exception:
        pass

    _replace_in_slide(
        prs.slides[12],
        {
            "实验设置：两个任务 · 10 个数据集 · 16 个对比模型": (
                "实验设置：两个任务 · 10 个数据集 · Table 1 对比方法"
            ),
            "16 个对比基线 · 三类": "对比基线 · 三类（正文称 14；以 Table 1 为准）",
            "Sundial · Chronos（含其他基础模型，共 16 个基线）": (
                "Sundial · Chronos 等（完整名单见 Table 1）"
            ),
        },
    )
    try:
        nt = prs.slides[12].notes_slide.notes_text_frame
        nt.text = nt.text.replace("基线一共 16 个", "基线以 Table 1 为准（正文写 fourteen）").replace(
            "16 个基线", "Table 1 基线"
        )
        if "本区数字" not in nt.text:
            nt.text += "\n【边界】本区数字均摘自论文；复现页另计。"
    except Exception:
        pass
    # footer on experiment slide
    for sh in prs.slides[12].shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "AlphaCast · 论文精读与组会分享":
            _set_shape_text(sh, "AlphaCast · 论文精读 · 本区数字摘自论文")

    _replace_in_slide(
        prs.slides[13],
        {
            "MSE 最优：9 / 10 数据集\nMAE 最优：10 / 10 数据集": (
                "MSE 最优：9 / 10 数据集\nMAE 最优：9 / 10 数据集"
            ),
            "例外：NP 市场 Chronos 的 MSE 22.180\n低于 AlphaCast 27.113——并非所有\n数据集全胜（如实呈现）": (
                "例外：NP MSE→Chronos 22.180（AC 27.113）；\n"
                "MOPEX MAE→TimeXer 1.277（AC 1.353）\n"
                "——并非所有数据集全胜"
            ),
            "数值来源：AlphaCast 表 1": "数值来源：论文表 1（不含本机复现）",
        },
    )
    try:
        nt = prs.slides[13].notes_slide.notes_text_frame
        nt.text = (
            nt.text.replace("MAE 上 10 个全部最优", "MAE 上 9 个最优（MOPEX 例外）")
            .replace("9/10 MSE 最优、10/10 MAE 最优", "MSE/MAE 均为 9/10 最优")
            .replace("10 / 10", "9 / 10")
        )
        if "MOPEX" not in nt.text:
            nt.text += "\n例外补充：MOPEX MAE 上 TimeXer 优于 AlphaCast。"
    except Exception:
        pass

    _replace_in_slide(
        prs.slides[17],
        {
            "5. 复现性：知识库构建 / 案例库维护等细节未完整说明": (
                "5. 复现性：知识库/案例库细节未写全；"
                "本机另有对齐工作（见复现页），不宣称 Table 1 系统格"
            ),
        },
    )

    # --- add two repro slides by duplicating blank-ish evaluation slide layout ---
    # Duplicate slide 17 (evaluation) twice at end, then fill and reorder
    eval_idx = 17  # 0-based current "我的评价"
    _duplicate_slide(prs, eval_idx)
    _duplicate_slide(prs, eval_idx)
    # Now slides: 0..19 original, 20=dup1, 21=dup2  (22 total)
    # Fill last two
    slides = list(prs.slides)
    _fill_repro_progress(slides[-2])
    _fill_repro_issues(slides[-1])

    # Reorder: 0..16 (through 深入分析), then repro pair (20,21), then 17,18,19
    # Current indices after append: 0..19 original, 20 progress, 21 issues
    order = list(range(17)) + [20, 21] + [17, 18, 19]
    _reorder_slides(prs, order)

    _update_page_numbers(prs)

    prs.save(str(PPTX))
    print(f"[ok] patched {PPTX} ({len(list(prs.slides))} slides)")


if __name__ == "__main__":
    patch()
