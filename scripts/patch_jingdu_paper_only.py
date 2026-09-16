#!/usr/bin/env python3
"""Patch 组会分享_AlphaCast_精读.pptx: paper-only + Investigator fact fix."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "docs" / "组会分享_AlphaCast_精读.pptx"


def _set_shape_text(shape, new_text: str) -> None:
    tf = shape.text_frame
    paras = list(tf.paragraphs)
    lines = new_text.split("\n")
    while len(paras) < len(lines):
        tf.add_paragraph()
        paras = list(tf.paragraphs)
    for i, para in enumerate(paras):
        line = lines[i] if i < len(lines) else ""
        if para.runs:
            para.runs[0].text = line
            for run in para.runs[1:]:
                run.text = ""
        else:
            run = para.add_run()
            run.text = line
    for i in range(len(lines), len(paras)):
        if paras[i].runs:
            paras[i].runs[0].text = ""
            for run in paras[i].runs[1:]:
                run.text = ""


def _replace_in_slide(slide, mapping: dict[str, str]) -> int:
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        old = sh.text_frame.text
        new = old
        for a, b in mapping.items():
            if a in new:
                new = new.replace(a, b)
        if new != old:
            _set_shape_text(sh, new)
            n += 1
    try:
        notes = slide.notes_slide.notes_text_frame
        old = notes.text
        new = old
        for a, b in mapping.items():
            if a in new:
                new = new.replace(a, b)
        if new != old:
            notes.text = new
            n += 1
    except Exception:
        pass
    return n


def _replace_exact_shape(slide, exact: str, new_text: str) -> bool:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == exact.strip():
            _set_shape_text(sh, new_text)
            return True
    # allow substring-only match on single-line boxes
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == exact or (exact in t and len(t) < len(exact) + 8):
            _set_shape_text(sh, new_text)
            return True
    return False


def _scrub_notes_repro(slide) -> None:
    try:
        nt = slide.notes_slide.notes_text_frame
        text = nt.text
        # drop sentences that are clearly local-repro coaching
        drop_pat = re.compile(
            r"[^\n]*(本机|CastMind|复现页|对齐等级|PAPER_ALIGNMENT|冒充 Table)[^\n]*\n?",
            re.M,
        )
        cleaned = drop_pat.sub("", text)
        if cleaned != text:
            nt.text = cleaned.strip() + ("\n" if cleaned.strip() else "")
    except Exception:
        pass


def _update_page_numbers(prs: Presentation) -> None:
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides):
        want = f"{i + 1:02d} / {total:02d}"
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if re.fullmatch(r"\d{1,2}\s*/\s*\d{1,2}", t):
                _set_shape_text(sh, want)


def patch() -> None:
    if not PPTX.is_file():
        raise SystemExit(f"missing {PPTX}")
    prs = Presentation(str(PPTX))
    assert len(prs.slides) >= 18, len(prs.slides)

    # --- Slide 1: strip repro footer ---
    _replace_in_slide(
        prs.slides[0],
        {
            "论文：AlphaCast（Zhang et al., USTC）· 会议 / arXiv 信息论文未标注 · 论文区数字均摘自论文表/图；复现页另计、不冒充 Table 1": (
                "论文：AlphaCast（Zhang et al., USTC）· 数字均摘自论文表/图（analysis/paper_facts 已核对）"
            ),
            "论文区数字均摘自论文表/图；复现页另计、不冒充 Table 1": (
                "论文区数字均摘自论文表/图"
            ),
            "复现页另计、不冒充 Table 1": "数字均摘自论文表/图",
        },
    )

    # --- Slide 2: route without local repro ---
    _replace_in_slide(
        prs.slides[1],
        {
            "汇报路线 · 六个问题（含本机复现边界）": "汇报路线 · 五个问题（纯论文精读）",
            "六个问题（含本机复现边界）": "五个问题（纯论文精读）",
            "价值判断 · 局限 · 本机复现边界": "价值判断 · 局限 · 开放问题",
            "本机复现边界": "开放问题",
        },
    )

    # --- Slide 3: TimesFM not in Table 1 ---
    _replace_in_slide(
        prs.slides[2],
        {
            "Chronos · Sundial · TimesFM（大规模预训练、少样本泛化）": (
                "Chronos · Sundial（§2.1；TimesFM 见附录，未进 Table 1）"
            ),
        },
    )

    # --- Slide 6–8: Investigator / feature wording (newline boxes) ---
    _replace_exact_shape(prs.slides[5], "相关证据", "F_selected")
    _replace_exact_shape(
        prs.slides[6],
        "调查者\n预处理 · 选择相关信息",
        "调查者\n语义解析 · 选出 F_selected",
    )
    _replace_exact_shape(
        prs.slides[7],
        "统计特征：均值 / 标准差 / 计数 / 自相关 / 季节强度 / 熵\n时间特征：刻画序列整体形态与周期性",
        "论文：约 20 个统计/时序特征（附录 B）\n案例分析常见 basic mean/std/count 等（§4.5.1）",
    )
    _replace_in_slide(
        prs.slides[7],
        {"生成者到底在看什么？": "四类信息源：材料从哪来？（生成者公式 8 综合使用）"},
    )

    # --- Slide 9: Investigator flowchart — NOT four sources ---
    s9 = prs.slides[8]
    _replace_exact_shape(s9, "全部信息（四个信息源）", "任务输入 I_input + 特征集 F")
    _replace_exact_shape(s9, "筛选出的相关证据", "选出的特征 F_selected")
    _replace_in_slide(
        s9,
        {
            "全部信息（四个信息源）": "任务输入 I_input + 特征集 F",
            "筛选出的相关证据": "选出的特征 F_selected",
            "F_selected → 交给生成者": (
                "F_selected → 交给生成者\n（K / E / 案例辅助与近邻由生成者综合，§3.4.1）"
            ),
            "根据任务提示判断本次预测需要哪些类型的信息": (
                "语义解析与需求分解：推断所需特征类型与背景（§3.4.1）"
            ),
            "从特征集筛出 F_selected = S(F, I_input)": (
                "特征选择：F_selected = S(F, I_input)（式 6）"
            ),
            "决定「对当前任务哪些信息最值得关注」而非堆砌信息": (
                "论文明确：从特征集筛选；【未写明】是否同时检索案例库"
            ),
            "“调查者更像一个研究员：先调查问题， | 再决定生成者应该重点看什么。”": (
                "四类信息源 ≠ 全进调查者：调查者主做特征选择；多源综合在生成者（式 8）"
            ),
            "回应论文动机 03：模型需要决定 | 「应该看哪些信息」，再谈如何预测": (
                "回应动机 03：先决定看哪些特征，再生成预测；四源材料供生成者使用"
            ),
        },
    )
    try:
        s9.notes_slide.notes_text_frame.text = (
            "论文 §3.4.1：Investigator 从 feature set 选 F_selected。"
            "Case 检索与 auxiliary 写在 Generator。"
            "不要说调查者看到全部四个信息源。"
        )
    except Exception:
        pass

    # --- Slide 10: keep generator multi-source accurate ---
    _replace_in_slide(
        prs.slides[9],
        {
            "生成者不是「从案例库挑一个最佳模型直接用」，而是综合多个证据来源进行预测": (
                "生成者综合式 (8) 输入：I_input、F_selected、K、E、X_auxiliary、X_neighbor → X_raw + CoT"
            ),
        },
    )

    # --- Slide 11: Reflector wording without unstated "future truth" claim as paper text ---
    _replace_in_slide(
        prs.slides[10],
        {
            "注意：未来真实值尚未出现，反思者检查的不是「数字准不准」，而是预测的合理性与推理质量": (
                "论文：按任务要求与证据充分性、CoT 审计（幻觉/循环/跳步）（§3.4.2）"
                "【分析】测试时不宜用未来真值在线纠错，否则泄漏——非论文原句"
            ),
        },
    )

    # --- Slide 13: baselines count ---
    _replace_in_slide(
        prs.slides[12],
        {
            "对比基线 · 三类（正文称 14；以 Table 1 为准）": (
                "对比基线 · 三类（正文称 fourteen；Table 1 另含 TimeXer/Transformer 等，以表为准）"
            ),
            "Sundial · Chronos 等（完整名单见 Table 1）": (
                "Sundial · Chronos（§4.1）；TimesFM 未进 Table 1"
            ),
            "AlphaCast · 论文精读 · 本区数字摘自论文": "AlphaCast · 论文精读 · 数字摘自论文",
        },
    )

    # --- Slide 14 footer ---
    _replace_in_slide(
        prs.slides[13],
        {
            "AlphaCast · 论文精读与组会分享 · 数值来源：论文表 1（不含本机复现）": (
                "AlphaCast · 论文精读 · 数值来源：论文 Table 1"
            ),
            "（不含本机复现）": "",
        },
    )

    # --- Slide 18: remove local repro bullet ---
    _replace_in_slide(
        prs.slides[17],
        {
            "5. 复现性：知识库/案例库细节未写全；本机另有对齐工作（见复现页），不宣称 Table 1 系统格": (
                "5. 复现性：知识库/案例库构建与维护细节论文未写全（§3.3）；权重与配方公开程度有限"
            ),
            "本机另有对齐工作（见复现页），不宣称 Table 1 系统格": (
                "知识库/案例库细节论文未写全"
            ),
            "见复现页": "",
        },
    )

    # --- Slide 19 Q&A: fix Q2/Q5 to paper; complete Q8 ---
    _replace_in_slide(
        prs.slides[18],
        {
            "普通做法是「历史数据 → 让 LLM 直接输出预测」；AlphaCast 让 LLM 先调查多源证据、再生成、再反思，形成闭环": (
                "普通做法常是单次映射；AlphaCast 用调查→生成→反思闭环，并接入特征/知识/情境/案例工具集"
            ),
            "信息并非越多越好：无关特征引入噪声、稀释注意力、增加成本；调查者按任务筛选出最相关的 F_selected": (
                "论文：调查者从特征集筛 F_selected（式 6）；消融显示去掉特征库会升 MSE（表 2）"
            ),
            "案例用于快速检索和提供辅助预测，但生成者综合五类证据推理，辅助预测只是参考之一": (
                "案例库提供辅助预测与近邻；生成者按式 (8) 综合多源，不是只用历史上最优模型交差"
            ),
            "它不检查数字准不准（真实值未出现），而是检查预测级与推理级质量；不合理就反馈迭代": (
                "检查预测级可靠性/证据与 CoT 质量（§3.4.2）；不合理则反馈调查者迭代"
            ),
            "问 8\n你认为最大的局限是什么？（个人观点）": (
                "问 8\n最大局限是什么？（个人观点）\n"
                "知识/案例细节与成本量化不足；骨干敏感（表 3）；Human Wisdom 如何体现仍可讨论"
            ),
        },
    )
    # If Q8 is only a title shape, try fill nearby empty
    for sh in prs.slides[18].shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == "你认为最大的局限是什么？（个人观点）":
            _set_shape_text(
                sh,
                "最大局限是什么？（个人观点）\n"
                "知识/案例细节与成本量化不足；骨干敏感（表 3）；Human Wisdom 如何编码仍可讨论",
            )

    # --- Global scrub of remaining repro phrases ---
    global_map = {
        "本机复现边界": "开放问题",
        "不含本机复现": "摘自论文",
        "见复现页": "",
        "本机另有对齐工作": "",
        "本区数字摘自论文": "数字摘自论文",
        "论文精读 · 本区数字摘自论文": "论文精读 · 数字摘自论文",
    }
    for slide in prs.slides:
        _replace_in_slide(slide, global_map)
        _scrub_notes_repro(slide)

    _update_page_numbers(prs)
    prs.save(str(PPTX))
    print(f"[ok] saved {PPTX} slides={len(prs.slides)}")


if __name__ == "__main__":
    patch()
